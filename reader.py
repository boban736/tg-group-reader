"""Hourly Telegram group digest: read new messages, let an LLM pick what matters,
append it to a daily note in an Obsidian vault and ping Saved Messages about urgent items."""

import asyncio
import base64
import datetime as dt
import io
import json
import logging
import os
import re
import sys
from dataclasses import dataclass, field
from pathlib import Path

from dotenv import load_dotenv
from openai import OpenAI
from telethon import TelegramClient, functions
from telethon.tl.types import (
    ChannelParticipantsAdmins,
    DocumentAttributeFilename,
    MessageMediaDocument,
    MessageMediaPhoto,
)

ROOT = Path(__file__).resolve().parent
load_dotenv(ROOT / ".env")

log = logging.getLogger("reader")


def env(name: str, default: str | None = None) -> str:
    value = os.getenv(name, default)
    if value is None or value == "":
        sys.exit(f"Missing required setting {name} in .env")
    return value


API_ID = int(env("TG_API_ID"))
API_HASH = env("TG_API_HASH")
GROUP = env("TG_GROUP")  # @username, numeric id (-100...) or exact chat title
OPENAI_MODEL = env("OPENAI_MODEL")
OPENAI_API_KEY = env("OPENAI_API_KEY")
OPENAI_BASE_URL = os.getenv("OPENAI_BASE_URL") or None

VAULT_DIR = Path(env("VAULT_DIR")).expanduser()
NOTES_DIR = VAULT_DIR / os.getenv("VAULT_NOTES_SUBDIR", "Универ/Дайджест")
ATTACH_DIR = VAULT_DIR / os.getenv("VAULT_ATTACH_SUBDIR", "Универ/Дайджест/files")
NOTE_TAGS = [t.strip() for t in os.getenv("NOTE_TAGS", "универ,дайджест").split(",") if t.strip()]
SUBJECT_LINKS = os.getenv("SUBJECT_LINKS", "1") == "1"  # [[Предмет]] wikilinks
SAVE_FILES_FOR = set(os.getenv("SAVE_FILES_FOR", "deadline,material,schedule,announcement").split(","))
STATE_FILE = ROOT / "state.json"
SESSION_FILE = ROOT / "tg.session"
PROMPT_FILE = ROOT / "prompts" / "system.md"

FIRST_RUN_HOURS = int(os.getenv("FIRST_RUN_HOURS", "24"))
MAX_MESSAGES = int(os.getenv("MAX_MESSAGES", "500"))
BATCH_MESSAGES = int(os.getenv("BATCH_MESSAGES", "150"))
BATCH_IMAGES = int(os.getenv("BATCH_IMAGES", "10"))
MAX_FILE_MB = float(os.getenv("MAX_FILE_MB", "20"))
DOC_TEXT_CHARS = int(os.getenv("DOC_TEXT_CHARS", "12000"))
IMAGE_DETAIL = os.getenv("IMAGE_DETAIL", "high")  # low | high | auto
NOTIFY_URGENT = os.getenv("NOTIFY_URGENT", "1") == "1"
NOTIFY_TARGET = os.getenv("NOTIFY_TARGET", "me")  # "me" = Saved Messages

IMAGE_EXTS = {".jpg", ".jpeg", ".png", ".webp", ".gif"}
TEXT_EXTS = {".txt", ".md", ".csv"}

CATEGORY_TITLES = {
    "deadline": "Дедлайны, тесты, экзамены",
    "schedule": "Замены и изменения пар",
    "material": "Файлы и материалы",
    "announcement": "Объявления",
    "other": "Прочее полезное",
}


@dataclass
class Msg:
    id: int
    date: dt.datetime
    sender: str
    is_admin: bool
    topic: str | None
    text: str
    reply_to: str | None = None
    attachments: list[str] = field(default_factory=list)  # human-readable notes / extracted text
    images: list[str] = field(default_factory=list)  # data: URIs
    files: list[tuple[str, bytes]] = field(default_factory=list)  # (name, data) to save into vault if useful
    link: str = ""


# ---------- state ----------

def load_state() -> dict:
    if STATE_FILE.exists():
        return json.loads(STATE_FILE.read_text())
    return {}


def save_state(state: dict) -> None:
    tmp = STATE_FILE.with_suffix(".tmp")
    tmp.write_text(json.dumps(state, indent=2))
    tmp.replace(STATE_FILE)


# ---------- telegram ----------

async def resolve_chat(client: TelegramClient):
    try:
        ident = int(GROUP) if GROUP.lstrip("-").isdigit() else GROUP
        return await client.get_entity(ident)
    except ValueError:
        async for d in client.iter_dialogs():
            if d.name == GROUP:
                return d.entity
    sys.exit(f"Chat {GROUP!r} not found. Use @username, -100... id or the exact title.")


async def fetch_admin_ids(client, chat) -> set[int]:
    try:
        return {u.id async for u in client.iter_participants(chat, filter=ChannelParticipantsAdmins)}
    except Exception as e:  # listing admins can be forbidden
        log.warning("cannot list admins: %s", e)
        return set()


async def fetch_topics(client, chat) -> dict[int, str]:
    if not getattr(chat, "forum", False):
        return {}
    topics: dict[int, str] = {1: "General"}
    # The request moved from channels.* to messages.* in newer layers.
    for make in (
        lambda: functions.messages.GetForumTopicsRequest(
            peer=chat, offset_date=None, offset_id=0, offset_topic=0, limit=100),
        lambda: functions.channels.GetForumTopicsRequest(
            channel=chat, offset_date=None, offset_id=0, offset_topic=0, limit=100),
    ):
        try:
            res = await client(make())
            topics.update({t.id: t.title for t in res.topics if hasattr(t, "title")})
            return topics
        except AttributeError:
            continue
        except Exception as e:
            log.warning("cannot fetch forum topics: %s", e)
            return topics
    return topics


def topic_id(m) -> int | None:
    r = m.reply_to
    if r is None:
        return 1
    if getattr(r, "forum_topic", False):
        return r.reply_to_top_id or r.reply_to_msg_id
    return 1


def sender_name(m) -> str:
    s = m.sender
    if s is None:
        return m.post_author or "Аноним"
    if getattr(s, "title", None):
        return m.post_author or s.title
    return " ".join(x for x in (s.first_name, s.last_name) if x) or s.username or str(s.id)


def message_link(chat, msg_id: int) -> str:
    if getattr(chat, "username", None):
        return f"https://t.me/{chat.username}/{msg_id}"
    return f"https://t.me/c/{chat.id}/{msg_id}"


def doc_filename(doc) -> str:
    for a in doc.attributes:
        if isinstance(a, DocumentAttributeFilename):
            return a.file_name
    return ""


def extract_doc_text(name: str, data: bytes) -> str | None:
    ext = Path(name).suffix.lower()
    try:
        if ext == ".pdf":
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(data))
            return "\n".join((p.extract_text() or "") for p in reader.pages)
        if ext == ".docx":
            import docx
            return "\n".join(p.text for p in docx.Document(io.BytesIO(data)).paragraphs)
        if ext == ".pptx":
            from pptx import Presentation
            prs = Presentation(io.BytesIO(data))
            return "\n".join(
                sh.text_frame.text for s in prs.slides for sh in s.shapes if sh.has_text_frame)
        if ext in TEXT_EXTS:
            return data.decode("utf-8", errors="replace")
    except Exception as e:
        log.warning("cannot extract text from %s: %s", name, e)
    return None


def to_data_uri(data: bytes, mime: str) -> str:
    return f"data:{mime};base64,{base64.b64encode(data).decode()}"


async def process_media(client, m, out: Msg) -> None:
    media = m.media
    if isinstance(media, MessageMediaPhoto):
        data = await client.download_media(m, file=bytes)
        if data:
            out.images.append(to_data_uri(data, "image/jpeg"))
            out.files.append(("photo.jpg", data))
            out.attachments.append("[фото — приложено изображением]")
        return
    if not isinstance(media, MessageMediaDocument) or media.document is None:
        if media is not None:
            out.attachments.append(f"[вложение: {type(media).__name__}]")
        return

    doc = media.document
    name = doc_filename(doc) or "файл"
    size_mb = doc.size / 1_048_576
    mime = doc.mime_type or ""
    ext = Path(name).suffix.lower()
    header = f"[файл: {name}, {size_mb:.1f} MB]"

    if size_mb > MAX_FILE_MB:
        out.attachments.append(header + " (слишком большой, не скачан)")
        return
    if mime.startswith("image/") or ext in IMAGE_EXTS:
        data = await client.download_media(m, file=bytes)
        out.images.append(to_data_uri(data, mime or "image/jpeg"))
        out.files.append((name, data))
        out.attachments.append(header + " — приложено изображением")
        return
    if ext in {".pdf", ".docx", ".pptx", *TEXT_EXTS}:
        data = await client.download_media(m, file=bytes)
        out.files.append((name, data))
        text = extract_doc_text(name, data)
        if text and text.strip():
            snippet = text.strip()[:DOC_TEXT_CHARS]
            out.attachments.append(f"{header}\n<<<содержимое файла\n{snippet}\nсодержимое файла>>>")
            return
    out.files.append((name, await client.download_media(m, file=bytes)))
    out.attachments.append(header)


async def collect(client, chat, state: dict) -> tuple[list[Msg], int]:
    key = str(chat.id)
    last_id = state.get(key, {}).get("last_id", 0)
    admins = await fetch_admin_ids(client, chat)
    topics = await fetch_topics(client, chat)

    if last_id:
        raw = [m async for m in client.iter_messages(chat, min_id=last_id, limit=MAX_MESSAGES)]
    else:
        since = dt.datetime.now(dt.timezone.utc) - dt.timedelta(hours=FIRST_RUN_HOURS)
        raw = []
        async for m in client.iter_messages(chat, limit=MAX_MESSAGES):
            if m.date < since:
                break
            raw.append(m)
    raw.reverse()
    if not raw:
        return [], last_id

    by_id = {m.id: m for m in raw}
    missing = {
        m.reply_to.reply_to_msg_id for m in raw
        if m.reply_to and m.reply_to.reply_to_msg_id
        and m.reply_to.reply_to_msg_id != topic_id(m)
        and m.reply_to.reply_to_msg_id not in by_id
    }
    if missing:
        for r in await client.get_messages(chat, ids=list(missing)):
            if r:
                by_id.setdefault(r.id, r)

    msgs: list[Msg] = []
    for m in raw:
        if m.action is not None:  # service messages: joins, pins, topic edits
            continue
        tid = topic_id(m)
        out = Msg(
            id=m.id,
            date=m.date.astimezone(),
            sender=sender_name(m),
            is_admin=(m.sender_id in admins) or bool(m.post_author),
            topic=topics.get(tid) if topics else None,
            text=m.message or "",
            link=message_link(chat, m.id),
        )
        rid = m.reply_to.reply_to_msg_id if m.reply_to else None
        if rid and rid != tid and rid in by_id:
            parent = by_id[rid]
            ptext = (parent.message or ("[медиа]" if parent.media else "")).replace("\n", " ")
            out.reply_to = f"{sender_name(parent)}: {ptext[:200]}"
        try:
            await process_media(client, m, out)
        except Exception as e:
            log.warning("media of msg %s failed: %s", m.id, e)
            out.attachments.append("[вложение не удалось скачать]")
        msgs.append(out)
    return msgs, raw[-1].id


async def pinned_text(client, chat) -> str | None:
    try:
        full = await client(functions.channels.GetFullChannelRequest(chat))
        pid = full.full_chat.pinned_msg_id
        if pid:
            m = await client.get_messages(chat, ids=pid)
            return m.message if m else None
    except Exception:
        pass
    return None


# ---------- LLM ----------

def render_msg(m: Msg) -> str:
    head = f"#{m.id} {m.date:%d.%m %H:%M} | {m.sender}{' [admin]' if m.is_admin else ''}"
    if m.topic:
        head += f" | тема: {m.topic}"
    lines = [head]
    if m.reply_to:
        lines.append(f"  ↪ в ответ на {m.reply_to}")
    if m.text:
        lines.append(m.text)
    lines.extend(m.attachments)
    return "\n".join(lines)


def batches(msgs: list[Msg]):
    cur: list[Msg] = []
    imgs = 0
    for m in msgs:
        if cur and (len(cur) >= BATCH_MESSAGES or imgs + len(m.images) > BATCH_IMAGES):
            yield cur
            cur, imgs = [], 0
        cur.append(m)
        imgs += len(m.images)
    if cur:
        yield cur


def note_path(day: dt.date) -> Path:
    return NOTES_DIR / f"{day:%Y-%m-%d}.md"


def recent_digest_tail(chars: int = 4000) -> str:
    today = dt.date.today()
    text = "".join(
        p.read_text() for p in (note_path(today - dt.timedelta(days=1)), note_path(today)) if p.exists())
    return text[-chars:]


def ask_llm(llm: OpenAI, system: str, chat_title: str, pinned: str | None, batch: list[Msg]) -> list[dict]:
    now = dt.datetime.now().astimezone()
    intro = [
        f"Сейчас: {now:%A %d.%m.%Y %H:%M}.",
        f"Группа: {chat_title}.",
    ]
    if pinned:
        intro.append(f"Закреплённое сообщение: {pinned}")
    tail = recent_digest_tail()
    if tail:
        intro.append("Уже записано в дайджест ранее (не повторяй это, если нет новых деталей):\n" + tail)
    intro.append("Новые сообщения:\n\n" + "\n\n".join(render_msg(m) for m in batch))

    content: list[dict] = [{"type": "text", "text": "\n\n".join(intro)}]
    for m in batch:
        for uri in m.images:
            content.append({"type": "text", "text": f"Изображение из сообщения #{m.id}:"})
            content.append({"type": "image_url", "image_url": {"url": uri, "detail": IMAGE_DETAIL}})

    resp = llm.chat.completions.create(
        model=OPENAI_MODEL,
        messages=[{"role": "system", "content": system}, {"role": "user", "content": content}],
        response_format={"type": "json_object"},
    )
    raw = resp.choices[0].message.content or "{}"
    try:
        return json.loads(raw).get("items", [])
    except json.JSONDecodeError:
        log.error("LLM returned invalid JSON: %s", raw[:500])
        raise


# ---------- output ----------

def safe_name(name: str) -> str:
    return re.sub(r'[\\/:*?"<>|#^\[\]]', "_", name).strip() or "file"


def save_files(items: list[dict], msgs: dict[int, Msg]) -> dict[int, list[str]]:
    """Save attachments of messages that useful items refer to; return msg id -> vault file names."""
    saved: dict[int, list[str]] = {}
    wanted = {mid for i in items if i.get("category", "other") in SAVE_FILES_FOR
              for mid in i.get("source_ids", [])}
    for mid in sorted(wanted):
        m = msgs.get(mid)
        if not m or not m.files:
            continue
        ATTACH_DIR.mkdir(parents=True, exist_ok=True)
        for n, (name, data) in enumerate(m.files):
            fname = safe_name(f"{m.date:%Y-%m-%d}_{m.id}_{n}_{name}" if name == "photo.jpg"
                              else f"{m.date:%Y-%m-%d}_{m.id}_{name}")
            (ATTACH_DIR / fname).write_bytes(data)
            saved.setdefault(mid, []).append(fname)
    return saved


def write_digest(items: list[dict], msgs: dict[int, Msg], chat_title: str) -> None:
    now = dt.datetime.now().astimezone()
    path = note_path(now.date())
    files = save_files(items, msgs)

    parts = []
    if not path.exists():
        tags = "".join(f"\n  - {t}" for t in NOTE_TAGS)
        parts.append(f"---\ndate: {now:%Y-%m-%d}\nsource: \"{chat_title}\"\ntags:{tags}\n---\n"
                     f"\n# Дайджест {now:%d.%m.%Y}\n")
    parts.append(f"\n## {now:%H:%M}\n")
    for cat, title in CATEGORY_TITLES.items():
        group = [i for i in items if i.get("category", "other") == cat]
        if not group:
            continue
        parts.append(f"\n### {title}\n")
        for i in group:
            mark = "🔴 " if i.get("urgent") else ""
            subj = i.get("subject")
            subj = f"[[{subj}]]: " if subj and SUBJECT_LINKS else (f"{subj}: " if subj else "")
            when = f" — **{i['when']}**" if i.get("when") else ""
            ids = [mid for mid in i.get("source_ids", []) if mid in msgs]
            refs = " ".join(f"[#{mid}]({msgs[mid].link})" for mid in ids)
            parts.append(f"- {mark}{subj}{i.get('summary', '').strip()}{when} {refs}".rstrip() + "\n")
            for mid in ids:
                for fname in files.get(mid, []):
                    embed = "!" if Path(fname).suffix.lower() in IMAGE_EXTS else ""
                    parts.append(f"\t- {embed}[[{fname}]]\n")

    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("a") as f:
        f.write("".join(parts))


async def notify(client, items: list[dict], chat_title: str, msgs: dict[int, Msg]) -> None:
    urgent = [i for i in items if i.get("urgent")]
    if not (NOTIFY_URGENT and urgent):
        return
    lines = [f"🔴 {chat_title}:"]
    for i in urgent:
        when = f" ({i['when']})" if i.get("when") else ""
        ref = next((msgs[mid].link for mid in i.get("source_ids", []) if mid in msgs), "")
        lines.append(f"• {i.get('summary', '').strip()}{when} {ref}".rstrip())
    await client.send_message(NOTIFY_TARGET, "\n".join(lines), link_preview=False)


# ---------- main ----------

async def main() -> None:
    logging.basicConfig(
        level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s",
        handlers=[logging.StreamHandler(), logging.FileHandler(ROOT / "reader.log")],
    )
    system = PROMPT_FILE.read_text()
    llm = OpenAI(api_key=OPENAI_API_KEY, base_url=OPENAI_BASE_URL)
    state = load_state()

    async with TelegramClient(str(SESSION_FILE), API_ID, API_HASH) as client:
        chat = await resolve_chat(client)
        title = getattr(chat, "title", GROUP)
        msgs, new_last = await collect(client, chat, state)
        if not msgs:
            log.info("no new messages")
            if new_last:
                state[str(chat.id)] = {"last_id": new_last}
                save_state(state)
            return

        log.info("%d new messages, %d images", len(msgs), sum(len(m.images) for m in msgs))
        pinned = await pinned_text(client, chat)
        by_id = {m.id: m for m in msgs}
        items: list[dict] = []
        for b in batches(msgs):
            items.extend(ask_llm(llm, system, title, pinned, b))

        if items:
            write_digest(items, by_id, title)
            await notify(client, items, title, by_id)
        log.info("%d items written", len(items))
        # Advance only after a successful run so a failure is retried next hour.
        state[str(chat.id)] = {"last_id": new_last}
        save_state(state)


if __name__ == "__main__":
    asyncio.run(main())
